import io
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from fpdf import FPDF


TEMPLATE_PATH = Path("data1") / "国家重点研发计划项目预申报书.docx"
PDF_FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/msyh.ttc"),
    Path("C:/Windows/Fonts/simsun.ttc"),
]
DEFAULT_EAST_ASIA_FONT = "宋体"


@dataclass
class ApplicationPayload:
    company_info: str
    policy_info: str
    project_background: str = ""
    innovation_points: str = ""
    technical_route: str = ""
    expected_results: str = ""


def analyze_template_structure(doc: Document) -> Dict[str, List[Dict[str, Any]]]:
    structure: Dict[str, List[Dict[str, Any]]] = {"sections": [], "tables": []}
    current_section: Optional[Dict[str, Any]] = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        is_heading = bool(para.runs and para.runs[0].bold) or para.style.name.startswith("Heading")
        if is_heading:
            current_section = {"title": text, "content": []}
            structure["sections"].append(current_section)
        elif current_section:
            current_section["content"].append(text)

    for table in doc.tables:
        structure["tables"].append(
            {
                "rows": len(table.rows),
                "cols": len(table.columns),
                "headers": [cell.text.strip() for cell in table.rows[0].cells] if table.rows else [],
            }
        )

    return structure


def _extract_template_styles(template_doc: Document) -> Dict[str, Dict[str, Any]]:
    template_styles: Dict[str, Dict[str, Any]] = {}
    supported = ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]

    for style in template_doc.styles:
        if style.name not in supported:
            continue

        font_name = style.font.name or "Microsoft YaHei"
        font_size = style.font.size.pt if style.font.size else 12

        try:
            r_pr = style.font.element
            ea_font_elem = r_pr.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
            east_asia_font = (
                ea_font_elem.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', DEFAULT_EAST_ASIA_FONT)
                if ea_font_elem is not None
                else DEFAULT_EAST_ASIA_FONT
            )
        except Exception:
            east_asia_font = DEFAULT_EAST_ASIA_FONT

        template_styles[style.name] = {
            "font_name": font_name,
            "east_asia_font": east_asia_font,
            "font_size": font_size,
            "bold": style.font.bold,
            "italic": style.font.italic,
            "alignment": style.paragraph_format.alignment,
            "space_before": style.paragraph_format.space_before,
            "space_after": style.paragraph_format.space_after,
            "line_spacing": style.paragraph_format.line_spacing,
            "left_indent": style.paragraph_format.left_indent,
            "right_indent": style.paragraph_format.right_indent,
            "first_line_indent": style.paragraph_format.first_line_indent,
        }

    return template_styles


def _apply_template_styles(doc: Document, template_styles: Dict[str, Dict[str, Any]]) -> None:
    for style_name, style_info in template_styles.items():
        if style_name not in doc.styles:
            continue

        style = doc.styles[style_name]
        style.font.name = style_info["font_name"]
        style.font.size = Pt(style_info["font_size"])
        style.font.bold = style_info["bold"]
        style.font.italic = style_info["italic"]
        style.paragraph_format.alignment = style_info["alignment"]
        style.paragraph_format.space_before = style_info["space_before"]
        style.paragraph_format.space_after = style_info["space_after"]
        style.paragraph_format.line_spacing = style_info["line_spacing"]
        style.paragraph_format.left_indent = style_info["left_indent"]
        style.paragraph_format.right_indent = style_info["right_indent"]
        style.paragraph_format.first_line_indent = style_info["first_line_indent"]

        try:
            r_pr = style.font.element
            r_fonts = r_pr.find(qn("w:rFonts"))
            if r_fonts is None:
                r_fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, r_fonts)
            r_fonts.set(qn("w:eastAsia"), style_info["east_asia_font"])
            r_fonts.set(qn("w:hAnsi"), style_info["font_name"])
        except Exception:
            pass


def _sanitize_markdown_text(text: str) -> str:
    return text.replace("**", "").replace("__", "").replace("~~", "")


def _ensure_template_exists() -> None:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"未找到申报书模板：{TEMPLATE_PATH}")


def build_application_prompt(payload: ApplicationPayload, sections: List[Dict[str, Any]]) -> str:
    section_titles = "\n".join(f"- {section['title']}" for section in sections)
    if not section_titles:
        section_titles = "- 项目背景\n- 研究内容\n- 预期成果"

    return (
        "请根据以下信息，生成一份完整、正式、可用于项目申报的申报书内容。\n\n"
        f"企业信息：\n{payload.company_info}\n\n"
        f"政策信息：\n{payload.policy_info}\n\n"
        f"项目背景：\n{payload.project_background}\n\n"
        f"创新点：\n{payload.innovation_points}\n\n"
        f"技术路线：\n{payload.technical_route}\n\n"
        f"预期成果：\n{payload.expected_results}\n\n"
        f"请尽量按以下章节结构输出：\n{section_titles}\n\n"
        "写作要求：\n"
        "1. 全文使用中文。\n"
        "2. 语言正式、严谨，符合项目申报书书面表达习惯。\n"
        "3. 内容尽量具体，避免空泛套话。\n"
        "4. 突出项目必要性、创新性、可行性和预期效益。\n"
        "5. 使用 Markdown 标题结构组织内容，便于后续导出。"
    ).strip()


def generate_markdown(llm: Any, payload: ApplicationPayload) -> str:
    _ensure_template_exists()
    template_doc = Document(str(TEMPLATE_PATH))
    structure = analyze_template_structure(template_doc)
    prompt = build_application_prompt(payload, structure["sections"])
    response = llm.complete(prompt)
    text = getattr(response, "text", None) or str(response)
    return text.strip()


def markdown_to_word(markdown_text: str) -> io.BytesIO:
    _ensure_template_exists()
    template_doc = Document(str(TEMPLATE_PATH))
    template_styles = _extract_template_styles(template_doc)

    doc = Document()
    _apply_template_styles(doc, template_styles)

    for raw_line in markdown_text.splitlines():
        line = _sanitize_markdown_text(raw_line.strip())
        if not line:
            continue

        if line.startswith("# "):
            para = doc.add_paragraph(line[2:])
            if "Heading 1" in doc.styles:
                para.style = doc.styles["Heading 1"]
        elif line.startswith("## "):
            para = doc.add_paragraph(line[3:])
            if "Heading 2" in doc.styles:
                para.style = doc.styles["Heading 2"]
        elif line.startswith("### "):
            para = doc.add_paragraph(line[4:])
            if "Heading 3" in doc.styles:
                para.style = doc.styles["Heading 3"]
        elif line.startswith("- "):
            para = doc.add_paragraph()
            if "List Bullet" in doc.styles:
                para.style = doc.styles["List Bullet"]
            para.add_run(line[2:])
        else:
            para = doc.add_paragraph(line)
            if "Normal" in doc.styles:
                para.style = doc.styles["Normal"]
            for run in para.runs:
                try:
                    r_pr = run._r.get_or_add_rPr()
                    r_fonts = r_pr.find(qn("w:rFonts"))
                    if r_fonts is None:
                        r_fonts = OxmlElement("w:rFonts")
                        r_pr.insert(0, r_fonts)
                    east_asia_font = template_styles.get("Normal", {}).get("east_asia_font", DEFAULT_EAST_ASIA_FONT)
                    r_fonts.set(qn("w:eastAsia"), east_asia_font)
                except Exception:
                    pass

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def markdown_to_pdf(markdown_text: str) -> io.BytesIO:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    font_path = next((path for path in PDF_FONT_CANDIDATES if path.exists()), None)
    if font_path is None:
        raise FileNotFoundError("未找到可用的中文字体文件，无法导出 PDF。")

    pdf.add_font("AppFont", "", str(font_path))
    pdf.set_font("AppFont", size=12)
    content_width = pdf.w - pdf.l_margin - pdf.r_margin

    def write_block(text: str, line_height: int) -> None:
        pdf.multi_cell(content_width, line_height, text)

    for raw_line in markdown_text.splitlines():
        line = _sanitize_markdown_text(raw_line.strip())
        if not line:
            pdf.ln(4)
            continue

        if line.startswith("# "):
            pdf.set_font("AppFont", size=18)
            write_block(line[2:], 10)
            pdf.ln(2)
            pdf.set_font("AppFont", size=12)
        elif line.startswith("## "):
            pdf.set_font("AppFont", size=16)
            write_block(line[3:], 9)
            pdf.ln(1)
            pdf.set_font("AppFont", size=12)
        elif line.startswith("### "):
            pdf.set_font("AppFont", size=14)
            write_block(line[4:], 8)
            pdf.ln(1)
            pdf.set_font("AppFont", size=12)
        elif line.startswith("- "):
            write_block("• " + line[2:], 8)
        else:
            write_block(line, 8)

    output = pdf.output(dest="S")
    if isinstance(output, bytearray):
        output = bytes(output)
    elif isinstance(output, str):
        output = output.encode("latin-1")

    buffer = io.BytesIO(output)
    buffer.seek(0)
    return buffer
